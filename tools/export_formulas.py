"""Exporta las celdas (constantes y fórmulas) y los nombres del libro para el motor de fórmulas de la web.

Ajustes para que el motor (HyperFormula) calcule igual que Excel:
- INDIRECT(celda): en este libro la celda siempre contiene un rango fijo construido con ADDRESS
  ("$L$2069:$L$2097"), así que se sustituye por la referencia directa. Si la celda no contiene texto
  (p.ej. un VERDADERO, error del propio Excel que IFERROR convierte en 0) se sustituye por NA().
- INDIRECT("Mejoras_"&Sheele!$M$5): se convierte en una elección explícita entre las tablas Mejoras_*.
- Se quitan los prefijos _xlfn./_xlws. y la hoja NamedRangesList (solo documentación).

Uso: python tools/export_formulas.py   (plantilla 8.7.0 -> public/plantilla.json)
     python tools/export_formulas.py <libro.xlsm> <salida.json>
"""
import json
import re
import sys
import openpyxl
from openpyxl.formula import Tokenizer
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import range_boundaries
from openpyxl.worksheet.formula import ArrayFormula

SKIP_SHEETS = {"NamedRangesList"}
MEJORAS = ["Agua", "Aire", "Fuego", "Ilusión", "Luz", "Naturaleza", "Oscuridad", "Sheele", "Tierra"]


def outside_strings(formula, fn):
    """Aplica fn solo a los trozos de la fórmula que no están entre comillas."""
    parts = formula.split('"')
    return '"'.join(fn(p) if i % 2 == 0 else p for i, p in enumerate(parts))


# --- Estrechar rangos: VLOOKUP/HLOOKUP/INDEX con columna o fila fija solo dependen de esa columna o fila.
# HyperFormula calcula dependencias de antemano y vería ciclos que Excel no ve (Excel solo sigue lo que lee).

def split_call(tokens, i):
    """tokens[i] es 'FUNC('. Devuelve (lista de argumentos como listas de tokens, índice tras el ')')."""
    args, cur, depth = [], [], 0
    j = i + 1
    while j < len(tokens):
        t = tokens[j]
        if t.type == "FUNC" and t.subtype == "OPEN" or t.type == "PAREN" and t.subtype == "OPEN":
            depth += 1
        elif t.type == "FUNC" and t.subtype == "CLOSE" or t.type == "PAREN" and t.subtype == "CLOSE":
            if depth == 0:
                args.append(cur)
                return args, j + 1, t.value
            depth -= 1
        elif t.type == "SEP" and t.subtype == "ARG" and depth == 0:
            args.append(cur)
            cur = []
            j += 1
            continue
        cur.append(t)
        j += 1
    raise ValueError("paréntesis sin cerrar")


def text(tokens):
    return "".join(t.value for t in tokens)


def as_range(tokens, sheet, names):
    """Si el argumento es un rango rectangular (o un nombre que lo es) devuelve (hoja, c0, r0, c1, r1)."""
    toks = [t for t in tokens if t.type != "WHITE-SPACE"]
    if len(toks) != 1 or toks[0].type != "OPERAND" or toks[0].subtype != "RANGE":
        return None
    ref = toks[0].value
    if ref.lower() in names:
        ref = names[ref.lower()].lstrip("=")
    m = re.fullmatch(r"(?:('?)(.+?)\1!)?(\$?[A-Z]{1,3}\$?\d+)(?::(\$?[A-Z]{1,3}\$?\d+))?", ref)
    if not m:
        return None
    sh = m.group(2) or sheet
    a, b = m.group(3).replace("$", ""), (m.group(4) or m.group(3)).replace("$", "")
    c0, r0, _, _ = range_boundaries(a)
    c1, r1, _, _ = range_boundaries(b)
    return sh, c0, r0, c1, r1


def ref_text(sh, c0, r0, c1, r1):
    a = f"${get_column_letter(c0)}${r0}"
    b = f"${get_column_letter(c1)}${r1}"
    return f"'{sh}'!{a}" + ("" if (c0, r0) == (c1, r1) else f":{b}")


def as_int(tokens, sheet=None, consts=None):
    """Número entero fijo: un literal, o una celda que contiene una constante numérica (p.ej. $AT$40 = 11)."""
    toks = [t for t in tokens if t.type != "WHITE-SPACE"]
    if len(toks) != 1 or toks[0].type != "OPERAND":
        return None
    v = None
    if toks[0].subtype == "NUMBER":
        v = float(toks[0].value)
    elif toks[0].subtype == "RANGE" and consts is not None:
        m = re.fullmatch(r"(?:('?)(.+?)\1!)?\$?([A-Z]{1,3})\$?(\d+)", toks[0].value)
        if m:
            c = consts.get((m.group(2) or sheet, m.group(3) + m.group(4)))
            v = float(c) if isinstance(c, (int, float)) and not isinstance(c, bool) else None
    return int(v) if v is not None and v.is_integer() and v >= 1 else None


def narrow(formula, sheet, names, consts=None):
    try:
        tokens = Tokenizer(formula).items
    except Exception:
        return formula

    def walk(toks):
        out, i = [], 0
        while i < len(toks):
            t = toks[i]
            if t.type == "FUNC" and t.subtype == "OPEN":
                args, end, closer = split_call(toks, i)
                args = [walk(a) for a in args]
                if t.value.endswith("("):
                    out.append(rewrite(t.value[:-1].upper(), args, t.value))
                else:  # matriz constante {..}
                    out.append(t.value + ",".join(args) + closer)
                i = end
            else:
                v = t.value
                if t.type == "OPERAND" and t.subtype == "TEXT" and '""' in v[1:-1]:
                    parts = v[1:-1].replace('""', '"').split('"')
                    v = "(" + "&CHAR(34)&".join(f'"{p}"' for p in parts) + ")"
                elif t.type == "OPERAND" and t.subtype == "RANGE":
                    # nombre que apunta a un rango -> el rango (HyperFormula evalúa los nombres como fórmulas)
                    rng = as_range([t], sheet, {**names, **{k.split("!", 1)[1]: e for k, e in names.items()
                                                             if k.startswith(sheet.lower() + "!")}})
                    if rng and not re.fullmatch(r"(?:'?.+?'?!)?\$?[A-Z]{1,3}\$?\d+(:\$?[A-Z]{1,3}\$?\d+)?", v):
                        v = ref_text(*rng)
                out.append(v)
                i += 1
        return "".join(out)

    def rewrite(fn, args, opener):
        plain = opener + ",".join(args) + ")"
        if fn in ("VLOOKUP", "HLOOKUP") and len(args) in (3, 4):
            rng, k = parse_arg(args[1]), parse_int(args[2])
            aproximada = not (len(args) == 4 and args[3].strip().upper() in ("FALSE", "FALSE()", "0"))
            if rng and not k and aproximada:
                # columna/fila variable y búsqueda aproximada: el motor falla con textos -> MATCHAPPROX (o exacta si
                # las claves son textos desordenados, como hace Excel en la práctica)
                sh, c0, r0, c1, r1 = rng
                cells = ([(sh, f"{get_column_letter(c0)}{r}") for r in range(r0, r1 + 1)] if fn == "VLOOKUP"
                         else [(sh, f"{get_column_letter(c)}{r0}") for c in range(c0, c1 + 1)])
                keys = [x for x in ([consts.get(c) for c in cells] if consts else []) if x is not None]
                desordenados = keys and all(isinstance(x, str) for x in keys) and [x.lower() for x in keys] != sorted(x.lower() for x in keys)
                key = ref_text(sh, c0, r0, c0, r1) if fn == "VLOOKUP" else ref_text(sh, c0, r0, c1, r0)
                pos = f"MATCH({args[0]},{key},0)" if desordenados else f"MATCHAPPROX({args[0]},{key},1)"
                return f"INDEX({args[1]},{pos},{args[2]})" if fn == "VLOOKUP" else f"INDEX({args[1]},{args[2]},{pos})"
            if rng and k:
                sh, c0, r0, c1, r1 = rng
                exact = not aproximada
                if not exact and consts is not None:
                    # búsqueda aproximada sobre textos desordenados: Excel devuelve la coincidencia exacta
                    # (p.ej. tipos de Sheele "Aire, Agua, Fuego..."); se usa búsqueda exacta para reproducirlo
                    cells = ([(sh, f"{get_column_letter(c0)}{r}") for r in range(r0, r1 + 1)] if fn == "VLOOKUP"
                             else [(sh, f"{get_column_letter(c)}{r0}") for c in range(c0, c1 + 1)])
                    keys = [consts.get(x) for x in cells]
                    keys = [k for k in keys if k is not None]
                    if keys and all(isinstance(k, str) for k in keys) and [k.lower() for k in keys] != sorted(k.lower() for k in keys):
                        exact = True
                if fn == "VLOOKUP" and c0 + k - 1 <= c1:
                    key, val = ref_text(sh, c0, r0, c0, r1), ref_text(sh, c0 + k - 1, r0, c0 + k - 1, r1)
                elif fn == "HLOOKUP" and r0 + k - 1 <= r1:
                    key, val = ref_text(sh, c0, r0, c1, r0), ref_text(sh, c0, r0 + k - 1, c1, r0 + k - 1)
                else:
                    return plain
                pos = f"MATCH({args[0]},{key},0)" if exact else f"MATCHAPPROX({args[0]},{key},1)"
                return f"INDEX({val},{pos})" if fn == "VLOOKUP" else f"INDEX({val},1,{pos})"
        if fn in ("COLUMN", "ROW") and len(args) == 1 and args[0].strip().upper().startswith("INDEX("):
            # COLUMN(INDEX(R,f,c)): en Excel INDEX devuelve una referencia; en el motor un valor -> se calcula la posición
            try:
                toks = Tokenizer("=" + args[0].strip()).items
                inner, end, _ = split_call(toks, 0)
                if end == len(toks) and len(inner) in (2, 3):
                    rng = parse_arg(text(inner[0]))
                    if rng:
                        sh, c0, r0, c1, r1 = rng
                        una_fila = r0 == r1 and len(inner) == 2
                        f_, c_ = ("1", text(inner[1])) if una_fila else (text(inner[1]), text(inner[2]) if len(inner) == 3 else "1")
                        if fn == "COLUMN":
                            return f"({c0}+({c_ or '1'})-1)"
                        return f"({r0}+({f_ or '1'})-1)"
            except Exception:
                pass
        if fn == "MATCH" and (len(args) == 2 or len(args) == 3 and args[2].strip() in ("1", "-1")):
            # MATCH aproximado de HyperFormula falla con textos: función propia (src/engine)
            return f"MATCHAPPROX({','.join(args)})"
        if fn == "INDEX" and len(args) in (2, 3):
            rng = parse_arg(args[0])
            if rng:
                sh, c0, r0, c1, r1 = rng
                if len(args) == 2 and r0 == r1 and c1 > c0:  # Excel: INDEX(fila, n) es la columna n
                    return f"INDEX({args[0]},1,{args[1]})"
                r = parse_int(args[1])
                c = parse_int(args[2]) if len(args) == 3 else None
                row_empty = args[1].strip() == ""
                if len(args) == 3 and row_empty and c and c0 + c - 1 <= c1:      # INDEX(R,,k) -> columna k
                    return ref_text(sh, c0 + c - 1, r0, c0 + c - 1, r1)
                if r and c and r0 + r - 1 <= r1 and c0 + c - 1 <= c1:            # INDEX(R,f,k) -> una celda
                    return ref_text(sh, c0 + c - 1, r0 + r - 1, c0 + c - 1, r0 + r - 1)
                if len(args) == 3 and c and c0 + c - 1 <= c1 and not row_empty:  # INDEX(R,expr,k) -> INDEX(col,expr)
                    return f"INDEX({ref_text(sh, c0 + c - 1, r0, c0 + c - 1, r1)},{args[1]})"
                if len(args) == 3 and r and r0 + r - 1 <= r1 and args[2].strip():  # INDEX(R,f,expr) -> INDEX(fila,1,expr)
                    return f"INDEX({ref_text(sh, c0, r0 + r - 1, c1, r0 + r - 1)},1,{args[2]})"
        return plain

    def parse_arg(s):
        try:
            return as_range(Tokenizer("=" + s).items, sheet, names)
        except Exception:
            return None

    def parse_int(s):
        try:
            return as_int(Tokenizer("=" + s).items, sheet, consts)
        except Exception:
            return None

    return "=" + walk(tokens)


def fix(formula, sheet, values, names=None, consts=None):
    formula = re.sub(r"_xlfn\.|_xlws\.", "", formula)
    if names is not None:
        formula = narrow(formula, sheet, names, consts)
    # HyperFormula solo entiende TRUE()/FALSE() como funciones
    formula = outside_strings(formula, lambda s: re.sub(r"(?<![A-Za-z0-9_.!$'])(TRUE|FALSE)(?![A-Za-z0-9_(])", r"\1()", s))

    def cell_ref(m):
        v = values[sheet][m.group(1)].value
        if isinstance(v, str) and re.fullmatch(r"\$?[A-Z]+\$?\d+(:\$?[A-Z]+\$?\d+)?", v):
            return f"'{sheet}'!{v}"
        return "NA()"
    formula = re.sub(r"INDIRECT\(([A-Z]+\d+)\)", cell_ref, formula)
    if 'INDIRECT("Mejoras_"&Sheele!$M$5)' in formula:
        # M47:M57 piden las 11 primeras mejoras; Mejoras_Sheele (M29:M57) las incluiría a ellas mismas
        lista = lambda m: "'Tablas Sheele'!$M$29:$M$39" if m == "Sheele" else f"Mejoras_{m}"
        choice = "".join(f'IF(Sheele!$M$5="{m}",{lista(m)},' for m in MEJORAS) + "NA()" + ")" * len(MEJORAS)
        formula = formula.replace('INDIRECT("Mejoras_"&Sheele!$M$5)', choice)
    return formula


def adaptar_lista(formula, sheet, names, consts):
    """Fórmula de un desplegable para el motor: mismas adaptaciones que las celdas, salvo INDIRECT,
    que se resuelve al pedir la lista (depende de los datos de la ficha)."""
    f = re.sub(r"_xlfn\.|_xlws\.", "", formula)
    f = narrow("=" + f, sheet, names, consts)[1:]
    return outside_strings(f, lambda s: re.sub(r"(?<![A-Za-z0-9_.!$'])(TRUE|FALSE)(?![A-Za-z0-9_(])", r"\1()", s))


def nombres_y_constantes(wb):
    """Rangos con nombre (para estrecharlos) y constantes de celda (para índices fijos)."""
    ranges = {n.lower(): d.attr_text for n, d in wb.defined_names.items()}
    ranges.update({f"{ws.title}!{n}".lower(): d.attr_text for ws in wb.worksheets for n, d in ws.defined_names.items()})
    consts = {(ws.title, c.coordinate): c.value for ws in wb.worksheets for row in ws.iter_rows() for c in row
              if isinstance(c.value, (int, float, str)) and not isinstance(c.value, bool) and not str(c.value).startswith("=")}
    return ranges, consts


def export(path):
    wb = openpyxl.load_workbook(path)
    values = openpyxl.load_workbook(path, data_only=True)
    ranges, consts = nombres_y_constantes(wb)
    sheets = {}
    for ws in wb.worksheets:
        if ws.title in SKIP_SHEETS:
            continue
        cells = {}
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if v is None:
                    continue
                if isinstance(v, ArrayFormula):
                    v = v.text
                if not isinstance(v, (str, int, float, bool)):
                    v = str(v)
                if isinstance(v, str) and v.startswith("="):
                    v = fix(v, ws.title, values, ranges, consts)
                elif isinstance(v, str):
                    v = "'" + v  # texto literal: sin esto HyperFormula convierte "1." en número
                cells[c.coordinate] = v
        sheets[ws.title] = cells
    names = {n: "=" + re.sub(r"_xlfn\.|_xlws\.", "", d.attr_text) for n, d in wb.defined_names.items()}
    for ws in wb.worksheets:  # nombres de ámbito de hoja
        for n, d in ws.defined_names.items():
            names[f"{ws.title}!{n}"] = "=" + d.attr_text
    return {"sheets": sheets, "names": names}


if __name__ == "__main__":
    if len(sys.argv) == 1:
        from migrate import ROOT, TEMPLATE
        sys.argv += [TEMPLATE, f"{ROOT}/public/plantilla.json"]
    json.dump(export(sys.argv[1]), open(sys.argv[2], "w", encoding="utf-8"), ensure_ascii=False, separators=(",", ":"))
