"""Migra fichas de versiones antiguas del Excel a la plantilla v8.7.0 y las recalcula con Excel.

Copia solo las celdas de entrada (desbloqueadas). Las filas se emparejan por sus etiquetas
(celdas de texto bloqueadas), así sobrevive a filas insertadas entre versiones.
Uso: python tools/migrate.py   -> golden/<nombre>.xlsm + informe por consola
"""
import difflib
import os
import shutil
import sys
import openpyxl
from openpyxl.cell.cell import MergedCell

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DOCS = r"D:\Escritorio\dissidia"
TEMPLATE = os.path.join(DOCS, r"documentacion\Ficha Anima v8.7.0.xlsm")
FICHAS = {
    "sesshomaru": os.path.join(DOCS, r"aa-Sesshomaru\Sesshomaru 4.xlsm"),
    "lock": os.path.join(DOCS, r"a-Lock\Ficha_lock lvl 6.xlsm"),
    "ayane": os.path.join(DOCS, r"Pilars of reborn\Ayane akame lvl4.xlsm"),
}
OUT = os.path.join(ROOT, "golden")
# Hojas con entradas del jugador (las "Tablas*" son reglamento, no se copian)
SHEETS = ["Principal", "General", "PDs", "Combate", "Ki", "Creación de Técnicas", "Místicos",
          "Metamagia", "Sheele", "Psíquicos", "Elan", "Personalización", "Grimorio Magia",
          "Grimorio de Vía", "Grimorio Psíquica", "Resumen"]

# Bloques que se movieron entre versiones y el emparejamiento por filas no resuelve:
# versión -> hoja -> [(filas antiguas, columnas antiguas, desplazamiento filas, desplazamiento columnas)]
# Las opciones de impresión de Resumen no se migran (no afectan a cálculos).
_PRE_86 = {
    "PDs": [(range(5, 14), range(16, 17), 2, -1),   # categorías 1-5: P5.. -> O7..
            (range(5, 14), range(19, 20), 2, 0)],   # niveles:        S5.. -> S7..
    "Sheele": [(range(5, 6), range(4, 7), 0, 9)],   # tipo de Sheele, vinculada: D5,F5 -> M5,O5
}
_PRE_85 = {**_PRE_86, "Personalización": [(range(23, 60), range(3, 18), 5, 0)]}  # bloque izquierdo (C:Q) +5 filas
# 8.4.1 y 8.4.2: el bloque derecho de Personalización (V:AF) está 2 filas más arriba que en la 8.7.0; el emparejamiento por
# etiquetas de la columna izquierda deja huecos en las filas 76-79 (legados de sangre) y 123 (calidad de la 3.ª marioneta)
_PRE_843 = {**_PRE_85, "Personalización": _PRE_85["Personalización"] + [(range(76, 80), range(22, 33), 2, 0), (range(123, 124), range(22, 33), 2, 0)]}
MOVES = {
    # Sheele se movió en 8.6.3; PDs en 8.6.0; Personalización en 8.5.0 (comprobado con las plantillas de cada versión)
    "8.6.2": {"Sheele": _PRE_86["Sheele"]},
    "8.6.1": {"Sheele": _PRE_86["Sheele"]},
    "8.6.0": {"Sheele": _PRE_86["Sheele"]},
    "8.5.0": _PRE_86,
    "8.4.3": _PRE_85,
    "8.4.2": _PRE_843,
    "8.4.1": _PRE_843,
}

# Opciones de desplegable renombradas en 8.7.0
RENAMES = {"Si": "Sí", "Aumentar car. psíquicas": "Aumentar car. mentales"}


def moved(version, sheet, row, col):
    for rows, cols, dr, dc in MOVES.get(version, {}).get(sheet, []):
        if row in rows and col in cols:
            return row + dr, col + dc
    return None


def is_formula(v):
    return isinstance(v, str) and v.startswith("=")


def form_width(ws):
    """Columnas del formulario: todo lo que queda a la izquierda de 'Zona de tablas auxiliares'."""
    for c in ws[1]:
        if isinstance(c.value, str) and c.value.startswith("Zona de tablas auxiliares"):
            return c.column - 1
    return ws.max_column


def signature(ws, r, width):
    """Etiquetas fijas de una fila del formulario: texto en celdas bloqueadas."""
    return tuple(c.value.strip() for c in ws[r][:width]
                 if isinstance(c.value, str) and not is_formula(c.value)
                 and c.protection.locked is not False and c.value.strip())


def row_map(old, new):
    wa, wb = form_width(old), form_width(new)
    a = [signature(old, r, wa) for r in range(1, old.max_row + 1)]
    b = [signature(new, r, wb) for r in range(1, new.max_row + 1)]
    m = {}
    pa = pb = 0  # fin del bloque anterior
    for blk in difflib.SequenceMatcher(None, a, b, autojunk=False).get_matching_blocks():
        # hueco del mismo tamaño en ambas versiones = filas con etiqueta retocada, se emparejan 1:1
        if blk.a - pa == blk.b - pb:
            for i in range(blk.a - pa):
                m[pa + i + 1] = pb + i + 1
        for i in range(blk.size):
            m[blk.a + i + 1] = blk.b + i + 1
        pa, pb = blk.a + blk.size, blk.b + blk.size
    # resto: mismo desplazamiento que la fila emparejada anterior, si la etiqueta coincide
    used, off = set(m.values()), 0
    for r in range(1, len(a) + 1):
        if r in m:
            off = m[r] - r
        elif 0 < r + off <= len(b) and r + off not in used and a[r - 1] == b[r + off - 1]:
            m[r] = r + off
            used.add(r + off)
    return m


def plan(src_path, tpl):
    """Devuelve [(hoja, celda_destino, valor)] y una lista de avisos."""
    src = openpyxl.load_workbook(src_path)
    version = src["Resumen"]["AK105"].value.split()[-1]  # "Ficha Excel. Versión 8.5.0"
    writes, warnings = [], []
    for name in SHEETS:
        if name not in src.sheetnames:
            continue
        old, new = src[name], tpl[name]
        rm = row_map(old, new)
        for row in old.iter_rows():
            for c in row:
                if c.protection.locked is not False or isinstance(c, MergedCell):
                    continue
                if c.value in (None, "") or is_formula(c.value):
                    continue
                c.value = RENAMES.get(c.value, c.value)
                if name == "Resumen" and c.column >= 40:
                    continue  # AN+: opciones de impresión
                mv = moved(version, name, c.row, c.column)
                if mv:
                    writes.append((name, new.cell(*mv).coordinate, c.value))
                    continue
                r = rm.get(c.row)
                if r is not None and new.cell(r, c.column).value == c.value:
                    continue  # la plantilla ya tiene ese valor
                if r is None and new.cell(c.row, c.column).value == c.value:
                    continue
                if r is None:
                    warnings.append(f"{name}!{c.coordinate} sin fila destino: {c.value!r}")
                    continue
                dst = new.cell(r, c.column)
                if isinstance(dst, MergedCell) or dst.protection.locked is not False:
                    warnings.append(f"{name}!{c.coordinate}->{dst.coordinate} no es entrada: {c.value!r}")
                    continue
                writes.append((name, dst.coordinate, c.value))
    return writes, warnings


def excel_write(path, writes):
    import win32com.client
    xl = win32com.client.DispatchEx("Excel.Application")
    xl.Visible = False
    xl.DisplayAlerts = False
    xl.AutomationSecurity = 3  # macros desactivadas
    try:
        wb = xl.Workbooks.Open(path)
        xl.EnableEvents = False
        for sheet, coord, v in writes:
            rng = wb.Worksheets(sheet).Range(coord)
            if isinstance(v, str):
                rng.Formula = "'" + v  # texto literal, que Excel no lo convierta en fecha/número
            else:
                rng.Value = v
        xl.CalculateFull()
        invalid = []
        for sheet, coord, v in writes:
            try:
                if not wb.Worksheets(sheet).Range(coord).Validation.Value:
                    invalid.append(f"{sheet}!{coord}={v!r}")
            except Exception:
                pass  # celda sin validación
        wb.Save()
        wb.Close()
        return invalid
    finally:
        xl.Quit()


def main(names):
    os.makedirs(OUT, exist_ok=True)
    tpl = openpyxl.load_workbook(TEMPLATE)
    for name in names:
        writes, warnings = plan(FICHAS[name], tpl)
        dst = os.path.join(OUT, name + ".xlsm")
        shutil.copyfile(TEMPLATE, dst)
        invalid = excel_write(dst, writes)
        print(f"== {name}: {len(writes)} celdas copiadas, {len(warnings)} avisos")
        for w in warnings:
            print("   ", w)
        for w in invalid:
            print("    fuera de su desplegable en 8.7.0:", w)


if __name__ == "__main__":
    main(sys.argv[1:] or list(FICHAS))
