"""Extrae del Excel v8.7.0 lo que necesita la web.

  src/data/tablas.json   tablas del reglamento (rangos con nombre de las hojas "Tablas*")
  ref/inputs.json        inventario de celdas de entrada por hoja (etiqueta, valor por defecto, desplegable)
  ref/nombres.json       rangos con nombre de una celda en las hojas de ficha (nombres útiles para el esquema)
  golden/<ficha>.json    entradas y valores calculados de cada ficha de referencia

Uso: python tools/extract.py
"""
import json
import os
import openpyxl
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import range_boundaries
from migrate import FICHAS, OUT, SHEETS, TEMPLATE, form_width, is_formula

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def dump(path, data):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=1, default=str)
    print(f"{os.path.relpath(path, ROOT)}: {os.path.getsize(path) // 1024} KB")


def clean(v):
    if isinstance(v, float):
        v = round(v, 9)  # quita ruido de coma flotante de Excel (0.6000000000000001)
        return int(v) if v.is_integer() else v
    if isinstance(v, str):
        v = v.strip()
        return v if v else None
    return v


def named_ranges(wb):
    """[(nombre, hoja, (c0, r0, c1, r1))] de los rangos con nombre con destino rectangular."""
    out = []
    for name, dn in wb.defined_names.items():
        try:
            dests = list(dn.destinations)
        except Exception:
            continue
        for sheet, ref in dests:
            b = range_boundaries(ref.replace("$", ""))
            if None not in b and sheet in wb.sheetnames:
                out.append((name, sheet, b))
    return out


def tablas(wf, wv):
    out = {}
    for name, sheet, (c0, r0, c1, r1) in named_ranges(wf):
        if not sheet.startswith("Tablas") or (c0, r0) == (c1, r1):
            continue
        vs, fs = wv[sheet], wf[sheet]
        # cabecera = fila encima del rango; si falta, letra de columna
        cols, seen = [], set()
        for c in range(c0, c1 + 1):
            h = clean(vs.cell(r0 - 1, c).value) if r0 > 1 else None
            h = str(h) if isinstance(h, str) else get_column_letter(c)
            while h in seen:
                h += "_"
            seen.add(h)
            cols.append(h)
        rows, dynamic = [], []
        for i, r in enumerate(range(r0, r1 + 1)):
            rows.append([clean(vs.cell(r, c).value) for c in range(c0, c1 + 1)])
            if any(is_formula(fs.cell(r, c).value) for c in range(c0, c1 + 1)):
                dynamic.append(i)
        out[name] = {"hoja": sheet, "ref": f"{get_column_letter(c0)}{r0}:{get_column_letter(c1)}{r1}",
                     "columnas": cols, "filas": rows,
                     "dinamicas": dynamic}  # filas con fórmulas: dependen de la ficha, no son reglamento fijo
    return out


def label(ws, cell, width):
    """Texto fijo más cercano a la izquierda en la fila; si no hay, encima en la columna."""
    for c in range(cell.column - 1, 0, -1):
        v = ws.cell(cell.row, c).value
        if isinstance(v, str) and not is_formula(v) and v.strip():
            return v.strip()
    for r in range(cell.row - 1, max(cell.row - 6, 0), -1):
        v = ws.cell(r, cell.column).value
        if isinstance(v, str) and not is_formula(v) and v.strip():
            return v.strip()
    return None


def trasladar(formula, origen, destino):
    """La fórmula de un desplegable está escrita para la primera celda del rango: se ajusta a cada celda."""
    if formula is None or origen == destino:
        return formula
    from openpyxl.formula.translate import Translator
    try:
        return Translator("=" + formula, origin=origen).translate_formula(destino)[1:]
    except Exception:
        return formula


def validations(ws):
    out = {}
    for dv in ws.data_validations.dataValidation:
        # relativa a la 1ª celda del 1er rango; openpyxl reordena el sqref, así que se toma el rango de más arriba
        primero = min(dv.sqref.ranges, key=lambda g: (g.min_row, g.min_col))
        origen = f"{get_column_letter(primero.min_col)}{primero.min_row}"
        for rng in dv.sqref.ranges:
            for row in ws.iter_rows(min_row=rng.min_row, max_row=rng.max_row,
                                    min_col=rng.min_col, max_col=rng.max_col):
                for c in row:
                    out[c.coordinate] = trasladar(dv.formula1, origen, c.coordinate)
    return out


def validations_ext(path):
    """Desplegables en la extensión x14 del XML (openpyxl los descarta): {hoja: {celda: fórmula}}."""
    import re
    import zipfile
    from openpyxl.utils.cell import range_boundaries as rb
    z = zipfile.ZipFile(path)
    wbx = z.read("xl/workbook.xml").decode("utf-8")
    rels = z.read("xl/_rels/workbook.xml.rels").decode("utf-8")
    targets = dict(re.findall(r'<Relationship[^>]*Id="(rId\d+)"[^>]*Target="([^"]+)"', rels))
    targets.update({k: v for v, k in re.findall(r'<Relationship[^>]*Target="([^"]+)"[^>]*Id="(rId\d+)"', rels)})
    out = {}
    for name, rid in re.findall(r'<sheet name="([^"]+)" sheetId="\d+"[^>]*r:id="(rId\d+)"', wbx):
        name = name.replace("&amp;", "&")
        xml = z.read("xl/" + targets[rid].lstrip("/").removeprefix("xl/")).decode("utf-8")
        cells = {}
        for f, sq in re.findall(r"<x14:dataValidation .*?<xm:f>(.*?)</xm:f>.*?<xm:sqref>(.*?)</xm:sqref>", xml, re.S):
            f = f.replace("&amp;", "&").replace("&quot;", '"').replace("&lt;", "<").replace("&gt;", ">")
            origen = sq.split()[0].split(":")[0]  # la fórmula es relativa a la 1ª celda del 1er rango
            for rng in sq.split():
                c0, r0, c1, r1 = rb(rng)
                for r in range(r0, r1 + 1):
                    for c in range(c0, c1 + 1):
                        cells[f"{get_column_letter(c)}{r}"] = trasladar(f, origen, f"{get_column_letter(c)}{r}")
        out[name] = cells
    return out


def input_cells(ws):
    width = form_width(ws)
    return [c for row in ws.iter_rows(max_col=width) for c in row
            if c.protection.locked is False and not isinstance(c, MergedCell)]


def inventory(wf):
    out = {}
    ext = validations_ext(TEMPLATE)
    for sheet in SHEETS:
        ws = wf[sheet]
        dv = {**validations(ws), **ext.get(sheet, {})}
        out[sheet] = [{"celda": c.coordinate, "etiqueta": label(ws, c, 0),
                       "defecto": None if is_formula(c.value) else clean(c.value),
                       "formula": c.value if is_formula(c.value) else None,
                       "lista": dv.get(c.coordinate)} for c in input_cells(ws)]
        print(f"   {sheet}: {len(out[sheet])} entradas")
    return out


def nombres(wf):
    return {name: f"{sheet}!{get_column_letter(c0)}{r0}"
            for name, sheet, (c0, r0, c1, r1) in named_ranges(wf)
            if (c0, r0) == (c1, r1) and not sheet.startswith("Tablas")}


def golden(path, wf_tpl):
    wf = openpyxl.load_workbook(path)
    wv = openpyxl.load_workbook(path, data_only=True)
    inputs, values = {}, {}
    for sheet in SHEETS:
        for c in input_cells(wf[sheet]):
            raw = wv[sheet][c.coordinate].value
            v = raw if isinstance(raw, str) else clean(raw)  # textos tal cual: un " " también es una entrada
            if v not in (None, "") and not is_formula(c.value):
                inputs[f"{sheet}!{c.coordinate}"] = v
        width = form_width(wf[sheet])
        for row in wf[sheet].iter_rows(max_col=width):
            for c in row:
                if is_formula(c.value):
                    v = clean(wv[sheet][c.coordinate].value)
                    if v is not None:
                        values[f"{sheet}!{c.coordinate}"] = v
    return {"entradas": inputs, "valores": values}


def main():
    wf = openpyxl.load_workbook(TEMPLATE)
    wv = openpyxl.load_workbook(TEMPLATE, data_only=True)
    dump(os.path.join(ROOT, "src", "data", "tablas.json"), tablas(wf, wv))
    inv = inventory(wf)
    dump(os.path.join(ROOT, "ref", "inputs.json"), inv)
    # desplegables de cada celda de entrada, para la web: {"Hoja!Celda": "fórmula de la lista"}
    from export_formulas import adaptar_lista, nombres_y_constantes
    ranges, consts = nombres_y_constantes(wf)
    listas = {f"{h}!{x['celda']}": adaptar_lista(x["lista"], h, ranges, consts) for h, xs in inv.items() for x in xs if x["lista"]}
    with open(os.path.join(ROOT, "src", "data", "listas.json"), "w", encoding="utf-8") as f:
        json.dump(listas, f, ensure_ascii=False, separators=(",", ":"))
    print(f"src/data/listas.json: {len(listas)} desplegables")
    dump(os.path.join(ROOT, "ref", "nombres.json"), nombres(wf))
    for name in FICHAS:
        dump(os.path.join(OUT, name + ".json"), golden(os.path.join(OUT, name + ".xlsm"), wf))


if __name__ == "__main__":
    main()
