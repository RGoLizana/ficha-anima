"""Plantilla exacta del PDF que genera la macro ExportPDF del Excel (hoja Resumen), para la web.

- Fondo (marco gris, cajas, líneas): se copia tal cual del PDF de referencia que exporta el Excel.
  Excel dibuja cada celda como un rectángulo, y es idéntico en todas las fichas.
- Textos: cada celda con contenido de la hoja Resumen se coloca en su caja (combinaciones, fuente,
  alineación, ajuste de texto). Las fronteras de filas y columnas se calibran contra los bordes de
  los rectángulos del PDF, porque Excel no escala igual en pantalla que al imprimir.

Salidas (todo en puntos PDF, página A4 595.32 x 841.92):
  src/data/pdf-layout.json   {pagina: {fondo, logo, textos}}; un texto es fijo ("text") o sale de una celda ("ref")
  ref/pdf/web/<ficha>.html   vista previa con los valores de la ficha golden

Uso: python tools/pdf_layout.py
"""
import colorsys
import html
import json
import os
import re
import zipfile
import openpyxl
import pdfplumber
from openpyxl.styles.colors import COLOR_INDEX
from openpyxl.utils import range_boundaries
from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter
from migrate import FICHAS, OUT, ROOT, is_formula

SHEET = "Resumen"
PAGES = {"resumen": "MarcaInicioResumen:MarcaFinalResumen", "notas": "MarcaInicioNotas:MarcaFinalNotas"}
REF_PDF = os.path.join(ROOT, "ref", "pdf", "Sesshomaru.pdf")  # exportado con la macro de la 8.7.0
A4 = (595.32, 841.92)


def pdf_size(pt):
    """Excel imprime al 94 % y redondea el cuerpo de letra a múltiplos de 0,12 pt (10 -> 9,36; 8 -> 7,56)."""
    return round(round(pt * 0.94 / 0.12) * 0.12, 2)


def col_px(width):
    """Ancho de columna de Excel (caracteres; Arial 10 = 7 px por dígito) a píxeles de pantalla."""
    return int((256 * width + int(128 / 7)) / 256 * 7)


def theme_colors(path):
    xml = zipfile.ZipFile(path).read("xl/theme/theme1.xml").decode("utf-8")
    scheme = re.search(r"<a:clrScheme.*?</a:clrScheme>", xml, re.S).group(0)
    c = dict(re.findall(r"<a:(dk1|lt1|dk2|lt2|accent\d)>.*?(?:srgbClr val|lastClr)=\"([0-9A-Fa-f]{6})\"", scheme, re.S))
    return [c["lt1"], c["dk1"], c["lt2"], c["dk2"]] + [c[f"accent{i}"] for i in range(1, 7)]  # orden de Excel


def resolve(color, theme, default="#000000"):
    try:
        if color.type == "rgb" and isinstance(color.rgb, str) and color.rgb != "00000000":
            return "#" + color.rgb[-6:]
        if color.type == "indexed" and color.indexed not in (64, 65):
            return "#" + COLOR_INDEX[color.indexed][-6:]
        if color.type == "theme":
            rgb = theme[color.theme]
            h, l, s = colorsys.rgb_to_hls(*(int(rgb[i:i + 2], 16) / 255 for i in (0, 2, 4)))
            t = color.tint or 0
            l = l * (1 + t) if t < 0 else l * (1 - t) + t
            return "#%02x%02x%02x" % tuple(round(v * 255) for v in colorsys.hls_to_rgb(h, l, s))
    except Exception:
        pass
    return default


def calibrate(bounds_px, edges_pt, start_pt, scale):
    """Asigna a cada frontera de fila/columna (px de pantalla) su posición en el PDF (pt).
    Recorre las fronteras en orden y engancha cada una al borde de rectángulo más cercano a lo esperado;
    las que no tienen borde visible se interpolan entre sus vecinas enganchadas."""
    anchors = {0: start_pt}  # (si una frontera no engancha, la siguiente amplía la tolerancia con la distancia)
    last_i, last_pt = 0, start_pt
    for i in range(1, len(bounds_px)):
        guess = last_pt + (bounds_px[i] - bounds_px[last_i]) * scale
        near = min(edges_pt, key=lambda e: abs(e - guess))
        step = (bounds_px[i] - bounds_px[i - 1]) * scale
        drift = 0.01 * (bounds_px[i] - bounds_px[last_i]) * scale
        if abs(near - guess) <= max(0.6, min(1.5, step * 0.4)) + drift:
            anchors[i] = near
            last_i, last_pt = i, near
    keys = sorted(anchors)
    out = []
    for i in range(len(bounds_px)):
        lo = max(k for k in keys if k <= i)
        hi = min((k for k in keys if k >= i), default=None)
        if hi is None or hi == lo:
            out.append(round(anchors[lo] + (bounds_px[i] - bounds_px[lo]) * scale, 2))
        else:
            f = (bounds_px[i] - bounds_px[lo]) / (bounds_px[hi] - bounds_px[lo])
            out.append(round(anchors[lo] + f * (anchors[hi] - anchors[lo]), 2))
    return out


def gray(c):
    v = c[0] if isinstance(c, (list, tuple)) else c
    return "#%02x%02x%02x" % ((round(float(v) * 255),) * 3)


def build(xlsm, pdf_path):
    wb = openpyxl.load_workbook(xlsm)
    theme = theme_colors(xlsm)
    ws = wb[SHEET]
    widths = {i: d.width for d in ws.column_dimensions.values() if d.width for i in range(d.min, d.max + 1)}
    default_w = ws.sheet_format.defaultColWidth or 8.43
    default_h = ws.sheet_format.defaultRowHeight or 12.75
    merges = {(m.min_row, m.min_col): (m.max_row, m.max_col) for m in ws.merged_cells.ranges}
    inside = {(r, c) for (r0, c0), (r1, c1) in merges.items()
              for r in range(r0, r1 + 1) for c in range(c0, c1 + 1)} - set(merges)
    pdf = pdfplumber.open(pdf_path)
    pages = {}
    for n, (page, spec) in enumerate(PAGES.items()):
        a, b = spec.split(":")
        refs = [list(wb.defined_names[x].destinations)[0][1].replace("$", "") for x in (a, b)]
        c0, r0, c1, r1 = range_boundaries(f"{refs[0]}:{refs[1]}")
        xs = [0]
        for c in range(c0, c1 + 1):
            xs.append(xs[-1] + col_px(widths.get(c, default_w)))
        height = {r: default_h if ws.row_dimensions[r].height is None else ws.row_dimensions[r].height
                  for r in range(r0, r1 + 1)}
        if page == "resumen":  # lo que hace FillPDFResumen antes de exportar
            box = sum(height[r] for r in range(4, 16))  # caja del retrato Y4:AJ15 +10 %
            height[4] += box * 0.05
            height[12] += box * 0.025
            height[15] += box * 0.025
            height[r1] = 0.1  # MarcaFinalResumen
        ys = [0]
        for r in range(r0, r1 + 1):
            ys.append(ys[-1] + height[r] * 4 / 3)

        p = pdf.pages[n]
        rects = p.rects
        ex = sorted({round(v, 2) for rc in rects for v in (rc["x0"], rc["x1"])})
        ey = sorted({round(v, 2) for rc in rects for v in (rc["top"], rc["bottom"])})
        X = calibrate(xs, ex, min(ex), 0.742)
        Y = calibrate(ys, ey, min(ey), 0.69)
        Y = calibrate(ys, ey, min(ey), (Y[-1] - Y[0]) / ys[-1])

        fondo = [[round(rc["x0"], 2), round(rc["top"], 2), round(rc["width"], 2), round(rc["height"], 2),
                  gray(rc.get("non_stroking_color") or 0)] for rc in rects]
        logo = [[round(i["x0"], 2), round(i["top"], 2), round(i["width"], 2), round(i["height"], 2)] for i in p.images]

        texts = []
        for r in range(r0, r1 + 1):
            for c in range(c0, c1 + 1):
                cell = ws.cell(r, c)
                if (r, c) in inside or cell.value in (None, ""):
                    continue
                rr, cc = merges.get((r, c), (r, c))
                rr, cc = min(rr, r1), min(cc, c1)
                f, al = cell.font, cell.alignment
                wrap = bool(al.wrap_text)
                # sin ajuste de texto, Excel deja desbordar a la derecha sobre celdas vacías
                over = cc + 1
                if not wrap:
                    while over <= c1 and (r, over) not in inside and (r, over) not in merges and ws.cell(r, over).value in (None, ""):
                        over += 1
                t = {"x": X[c - c0], "y": Y[r - r0], "w": round(X[cc + 1 - c0] - X[c - c0], 2),
                     "h": round(Y[rr + 1 - r0] - Y[r - r0], 2), "max_w": round(X[over - c0] - X[c - c0], 2),
                     "size": pdf_size(f.sz or 10), "bold": bool(f.b), "italic": bool(f.i),
                     "color": resolve(f.color, theme), "h_align": al.horizontal or "general",
                     "v_align": al.vertical or "bottom", "wrap": wrap, "cell": cell.coordinate}
                if is_formula(cell.value):
                    t["ref"] = f"{SHEET}!{cell.coordinate}"
                else:
                    t["text"] = str(cell.value)
                texts.append(t)
        pages[page] = {"range": f"{get_column_letter(c0)}{r0}:{get_column_letter(c1)}{r1}",
                       "fondo": fondo, "logo": logo, "textos": texts}
    pdf.close()
    return pages


def fmt(v):
    if isinstance(v, bool):
        return "VERDADERO" if v else "FALSO"
    if isinstance(v, float):
        return f"{v:g}"
    return str(v)


def render(page, values, logo_src=None):
    """HTML de una página A4 (unidades pt) con los valores de una ficha."""
    out = [f'<div style="position:relative;width:{A4[0]}pt;height:{A4[1]}pt;background:#fff;overflow:hidden;'
           f'font-family:Arial,\'Liberation Sans\',Helvetica,sans-serif">']
    for x, y, w, h, col in page["fondo"]:
        out.append(f'<div style="position:absolute;left:{x}pt;top:{y}pt;width:{w}pt;height:{h}pt;background:{col}"></div>')
    for x, y, w, h in page["logo"]:
        inner = (f'<img src="{logo_src}" alt="Anima Beyond Fantasy" style="width:100%;height:100%;object-fit:fill">' if logo_src else
                 '<div style="height:100%;display:flex;align-items:center;justify-content:center;border:1px dashed #999;color:#666;font:italic 9pt Arial">[Logo Anima]</div>')
        out.append(f'<div style="position:absolute;left:{x}pt;top:{y}pt;width:{w}pt;height:{h}pt">{inner}</div>')
    for t in page["textos"]:
        v = t["text"] if "text" in t else values.get(t["ref"])
        if v in (None, ""):
            continue
        ha = t["h_align"]
        if ha == "general":
            ha = "right" if isinstance(v, (int, float)) and not isinstance(v, bool) else "left"
        ha = {"centerContinuous": "center", "justify": "left", "fill": "left", "distributed": "center"}.get(ha, ha)
        va = {"top": "flex-start", "center": "center", "bottom": "flex-end"}.get(t["v_align"], "flex-end")
        x, w = t["x"], t["w"]
        if not t["wrap"] and ha == "left":
            w = t["max_w"]
        size, bold = t["size"], t["bold"]
        if isinstance(v, list):  # texto con formato por tramos (notas): [{"t", "b", "pt"}]
            size, bold = pdf_size(10), False
            content = "".join(f'<span style="font-weight:{700 if r["b"] else 400};font-size:{pdf_size(r["pt"])}pt">'
                              f'{html.escape(r["t"])}</span>' for r in v)
        else:
            content = html.escape(fmt(v))
        out.append(f'<div style="position:absolute;left:{x}pt;top:{t["y"]}pt;width:{w}pt;height:{t["h"]}pt;display:flex;'
                   f'align-items:{va};justify-content:{ {"left": "flex-start", "right": "flex-end", "center": "center"}[ha] };'
                   f'overflow:{"hidden" if t["wrap"] else "visible"};font-size:{size}pt;line-height:1.28;letter-spacing:0.0025em;'
                   f'font-weight:{700 if bold else 400};font-style:{"italic" if t["italic"] else "normal"};color:{t["color"]};'
                   f'padding:0 1.5pt;box-sizing:border-box">'
                   f'<span style="white-space:{"pre-wrap" if t["wrap"] else "pre"};text-align:{ha};{"width:100%" if t["wrap"] else ""}">'
                   f'{content}</span></div>')
    out.append("</div>")
    return "".join(out)


NOTAS = ["Principal", "Poderes", "PDs", "EquipoCombate", "CapacidadesCombate", "Ki", "Misticos", "Sheele",
         "Psiquicos", "Personalizacion"]


def notas(wb):
    """Texto de la página de notas tal como lo compone la macro FillPDFNotes (títulos en negrita a 12 pt)."""
    def val(name):
        sheet, ref = list(wb.defined_names[name].destinations)[0]
        return wb[sheet][ref.replace("$", "")].value
    runs = []
    if val("MostrarIdiomas") and val("NotasIdiomas"):
        runs += [{"t": val("NotasIdiomasTitulo"), "b": True, "pt": 12}, {"t": ": " + str(val("NotasIdiomas")) + "\n\n", "b": False, "pt": 10}]
    if val("MostrarNotas"):
        for n in NOTAS:
            if val(f"Notas{n}"):
                runs += [{"t": val(f"Notas{n}Titulo"), "b": True, "pt": 12}, {"t": "\n" + str(val(f"Notas{n}")) + "\n\n", "b": False, "pt": 10}]
    return runs


def values_of(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    vals = {f"{SHEET}!{c.coordinate}": c.value for row in wb[SHEET].iter_rows() for c in row if c.value is not None}
    vals[f"{SHEET}!D112"] = notas(wb)  # celda NotasResumen
    return vals


def main():
    pages = build(os.path.join(OUT, "sesshomaru.xlsm"), REF_PDF)
    with open(os.path.join(ROOT, "src", "data", "pdf-layout.json"), "w", encoding="utf-8") as f:
        json.dump(pages, f, ensure_ascii=False)
    print({k: (p["range"], len(p["fondo"]), len(p["textos"])) for k, p in pages.items()})
    os.makedirs(os.path.join(ROOT, "ref", "pdf", "web"), exist_ok=True)
    for name in FICHAS:
        vals = values_of(os.path.join(OUT, name + ".xlsm"))
        body = "".join(f'<div class="page">{render(p, vals)}</div>' for p in pages.values())
        doc = ('<!doctype html><meta charset="utf-8"><title>' + name + '</title>'
               '<style>@page{size:A4;margin:0}body{margin:0;background:#888}'
               '.page{width:595.32pt;height:841.92pt;margin:0 auto 12px;break-after:page}'
               '@media print{body{background:#fff}.page{margin:0}}</style>' + body)
        with open(os.path.join(ROOT, "ref", "pdf", "web", name + ".html"), "w", encoding="utf-8") as f:
            f.write(doc)
    print("ref/pdf/web/*.html")


if __name__ == "__main__":
    main()
