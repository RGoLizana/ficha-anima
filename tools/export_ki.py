"""Reconstruye los árboles de habilidades de la hoja 'Ki' (habilidades del Ki y del Némesis) para dibujarlos en la web.

Cómo se deduce (la hoja dibuja cada árbol como un esquema sangrado con caracteres '├ └ │'):
  * Cada habilidad es una fila con casilla de compra (claves de src/data/listas.json): Q10..Q64 para el Ki (nombre en
    K..N, coste en P) e I43..I64 para el Némesis (nombre en C..G, coste en H). El coste sale de la fórmula del coste,
    =IF(Q13=0,20,"-"): la celda de coste muestra "-" cuando la habilidad ya se tiene (comprada o innata, p.ej. Ocultación
    del Ki para los D'Anjayni).
  * Profundidad = columna del primer texto que no es un trazo (K=0, L=1...), +1 si el texto empieza por '└'/'├'
    ("└ Mult. arcana" en N21 cuelga de "Mult. mayor" en N20). El padre de una fila es la fila anterior más cercana con
    una profundidad menos.
  * La primera fila de cada árbol (Uso del Ki, Uso del Némesis) va en su propio recuadro encima del resto: es la raíz y
    las habilidades de profundidad 0 cuelgan de ella.
  * Inhumanidad y Zen se pueden tener sin comprarlas por las habilidades esenciales: la columna B de su fila lo calcula
    (COUNTIF(HabilidadesEsenciales_Adquiridas,"Inhumanidad")...) y se guarda como "innato".
  * Contraste con el Core (pp. 103-105, «Requisitos:» de cada habilidad): CORE abajo. Si el Core confirma una arista se
    anota en "core"; las aristas de la raíz sin glifo que el Core confirma llevan fuente 'Core p.N'. Si el Core choca con
    la hoja manda la hoja y se anota en "notas". Las habilidades que no están en el Core solo tienen la hoja como fuente.
  Validación: las compras de las fichas de golden/ no se saltan el árbol (se imprime al final).

Salida: src/data/ki-arbol.json
  {ki|nemesis: {raiz: clave, nodos: [{c clave, n nombre, p coste, cp celda de coste, padre clave|null, nivel, fila,
                                      innato?: celda}],
                aristas: [{de, a, fuente: 'hoja'|'Core p.N', core?: 'Core p.N'}]},
   notas: [texto]}
Uso: python tools/export_ki.py
"""
import glob
import json
import os
import re
import warnings

import openpyxl

warnings.filterwarnings("ignore")
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TRAZOS = " │├└─ "

# Core Exxet, capítulo VIII (Dominio del Ki): habilidad -> (requisito, página)
CORE = {
    "Uso del Ki": (None, 103), "Control del Ki": ("Uso del Ki", 103), "Detección del Ki": ("Control del Ki", 103),
    "Erudición": ("Detección del Ki", 103), "Eliminación de peso": ("Uso del Ki", 103), "Levitación": ("Eliminación de peso", 103),
    "Movimiento de objetos": ("Levitación", 103), "Vuelo": ("Levitación", 104), "Extrusión de presencia": ("Uso del Ki", 104),
    "Armadura de energía": ("Extrusión de presencia", 104), "Extensión del aura al arma": ("Extrusión de presencia", 104),
    "Destrucción por Ki": ("Extrusión de presencia", 104), "Transmisión del Ki": ("Uso del Ki", 104),
    "Curación por Ki": ("Transmisión del Ki", 104), "Uso de la energía necesaria": ("Uso del Ki", 104),
    "Ocultación del Ki": ("Uso de la energía necesaria", 105), "Falsa muerte": ("Ocultación del Ki", 105),
    "Eliminación de necesidades": ("Uso de la energía necesaria", 105),
    "Eliminación de penalizadores": ("Uso de la energía necesaria", 105), "Recuperación": ("Eliminación de penalizadores", 105),
    "Aumento de características": ("Uso de la energía necesaria", 105), "Inhumanidad": ("Uso del Ki", 105),
    "Zen": ("Inhumanidad", 105),
}

ARBOLES = {  # nombre: (columna de compra, columna de coste, columnas del nombre, filas)
    "ki": ("Q", "P", "KLMN", (10, 64)),
    "nemesis": ("I", "H", "CDEFG", (43, 64)),
}


def texto(v):
    if v is None:
        return ""
    v = str(v)
    m = re.fullmatch(r'="(.*)"', v)   # M39 = ="Alcance incrementado"
    return m.group(1) if m else v


def arbol(ws, listas, compra, coste, cols, filas):
    nodos, pila = [], []   # pila: último nodo visto en cada profundidad
    for fila in range(filas[0], filas[1] + 1):
        c = f"Ki!{compra}{fila}"
        if c not in listas:
            continue
        for i, col in enumerate(cols):
            t = texto(ws[f"{col}{fila}"].value)
            if t.strip(TRAZOS):
                break
        nombre = t.strip(TRAZOS)
        nivel = i + (1 if t.lstrip("  ")[:1] in "├└" else 0) + 1
        f = str(ws[f"{coste}{fila}"].value)
        p = int(re.search(r',(\d+),"-"\)$', f).group(1))
        if not nodos:
            nivel = 0   # recuadro de la raíz
        pila[nivel:] = [c]
        n = {"c": c, "n": nombre, "p": p, "cp": f"Ki!{coste}{fila}", "padre": pila[nivel - 1] if nivel else None,
             "nivel": nivel, "fila": fila}
        b = str(ws[f"B{fila}"].value or "")
        if "HabilidadesEsenciales" in b and f'"{nombre}"' in b:
            n["innato"] = f"Ki!B{fila}"
        nodos.append(n)
    return nodos


def main():
    listas = json.load(open(os.path.join(ROOT, "src", "data", "listas.json"), encoding="utf-8"))
    ws = openpyxl.load_workbook(os.path.join(ROOT, "golden", "lock.xlsm"))["Ki"]
    salida, notas = {}, []
    for nombre, (compra, coste, cols, filas) in ARBOLES.items():
        nodos = arbol(ws, listas, compra, coste, cols, filas)
        por = {n["c"]: n for n in nodos}
        aristas = []
        for n in nodos:
            if not n["padre"]:
                continue
            padre = por[n["padre"]]["n"]
            a = {"de": n["padre"], "a": n["c"], "fuente": "hoja"}
            core = CORE.get(n["n"]) if nombre == "ki" else None
            if core and core[0] == padre:
                a["core"] = f"Core p.{core[1]}"
                if n["nivel"] == 1:   # cuelga de la raíz por estar en su recuadro, sin glifo: lo confirma el Core
                    a["fuente"] = a["core"]
            elif core:
                notas.append(f"{n['n']}: el Core (p.{core[1]}) pide {core[0]}, la hoja la pone bajo {padre}; manda la hoja.")
            aristas.append(a)
        if nombre == "ki":
            fuera = [n["n"] for n in nodos if n["n"] not in CORE]
            notas.append(f"No están en el Core (solo la hoja): {', '.join(fuera)}.")
        else:
            notas.append("El árbol del Némesis no está en el Core: todas sus aristas son de la hoja.")
        salida[nombre] = {"raiz": nodos[0]["c"], "nodos": nodos, "aristas": aristas}
        print(nombre, len(nodos), "nodos,", len(aristas), "aristas")
    salida["notas"] = notas
    ruta = os.path.join(ROOT, "src", "data", "ki-arbol.json")
    with open(ruta, "w", encoding="utf-8") as fh:
        json.dump(salida, fh, ensure_ascii=False, indent=1)
    for t in notas:
        print(" nota:", t)
    validar(salida)


def validar(salida):
    for ruta in sorted(glob.glob(os.path.join(ROOT, "golden", "*.xlsm"))):
        ws = openpyxl.load_workbook(ruta, data_only=True, read_only=True)["Ki"]
        celdas = {f"Ki!{c.coordinate}": c.value for fila in ws.iter_rows(min_row=10, max_row=64, max_col=17) for c in fila if hasattr(c, "coordinate")}
        mal = []
        for nombre in ARBOLES:
            nodos = salida[nombre]["nodos"]
            tiene = {n["c"] for n in nodos if celdas.get(n["c"]) not in (None, "") or celdas.get(n["cp"]) == "-"
                     or ("innato" in n and celdas.get(n["innato"]))}
            mal += [n["n"] for n in nodos if celdas.get(n["c"]) not in (None, "") and n["padre"] and n["padre"] not in tiene]
        print(" ", os.path.basename(ruta), "OK" if not mal else f"SE SALTA EL ÁRBOL: {mal}")


if __name__ == "__main__":
    main()
