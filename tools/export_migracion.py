"""Mapa de celdas de versiones antiguas de la ficha a la 8.7.0, para importar .xlsm en el navegador.

Para cada versión: emparejamiento de filas por hoja (migrate.row_map, en tramos [fila_antigua, fila_nueva, n])
y los bloques movidos a mano (migrate.MOVES). Lo usa src/import/xlsm.ts.
Uso: python tools/export_migracion.py  -> src/data/migracion.json
"""
import json
import os
import openpyxl
from migrate import DOCS, MOVES, RENAMES, SHEETS, TEMPLATE, row_map

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
# un libro de cada versión (las etiquetas fijas son las mismas en todas las fichas de una versión)
MUESTRAS = {
    "8.4.1": r"Partida lluis\Kal lvl 2.xlsm",
    "8.4.2": r"partidas\Dissidia\la cena\Varja.xlsm",
    "8.4.3": r"a-Lock\Ficha_lock lvl 6.xlsm",
    "8.5.0": r"documentacion\Ficha Anima v8.5.0.xlsm",
    "8.6.0": r"documentacion\Ficha Anima v8.6.0.xlsm",
    "8.6.1": r"documentacion\Ficha_Anima_v8.6.1.xlsm",
    "8.6.2": r"documentacion\Ficha Anima v8.6.2.xlsm",
    "8.6.3": r"documentacion\Ficha Anima v8.6.3.xlsm",
}


def tramos(m):
    out = []
    for a in sorted(m):
        if out and a == out[-1][0] + out[-1][2] and m[a] == out[-1][1] + out[-1][2]:
            out[-1][2] += 1
        else:
            out.append([a, m[a], 1])
    return out


ANCLAS_HOJAS = ["Principal", "General", "PDs", "Combate", "Místicos", "Psíquicos", "Personalización"]


def anclas(tpl):
    """Etiquetas fijas (texto constante, hasta la fila 140) repartidas por cada hoja de entrada de la 8.7.0. Sirven para avisar
    cuando una ficha 8.7.0 tiene otra disposición (filas o columnas añadidas por un gremio) y sus datos podrían desplazarse."""
    out = {}
    for h in ANCLAS_HOJAS:
        c = []
        for row in tpl[h].iter_rows(min_row=1, max_row=140, max_col=52):
            for x in row:
                v = x.value
                if isinstance(v, str) and not v.startswith("=") and 4 <= len(v.strip()) <= 30 and x.column <= 52:
                    c.append((x.row, x.coordinate, " ".join(v.split())))
        c.sort()
        paso = max(1, len(c) // 12)
        out[h] = {k: v for _, k, v in c[::paso][:12]}
    return out


def main():
    tpl = openpyxl.load_workbook(TEMPLATE)
    versiones = {}
    for v, rel in MUESTRAS.items():
        old = openpyxl.load_workbook(os.path.join(DOCS, rel))
        filas = {h: tramos(row_map(old[h], tpl[h])) for h in SHEETS if h in old.sheetnames}
        mover = {h: [[r.start, r.stop - 1, c.start, c.stop - 1, dr, dc] for r, c, dr, dc in bs]
                 for h, bs in MOVES.get(v, {}).items()}
        versiones[v] = {"filas": filas, "mover": mover}
        print(v, sum(len(t) for t in filas.values()), "tramos")
    # celdas de entrada de la 8.7.0 y sus valores por defecto (de ref/inputs.json, que genera extract.py)
    with open(os.path.join(ROOT, "ref", "inputs.json"), encoding="utf-8") as f:
        inv = json.load(f)
    entradas = {h: " ".join(x["celda"] for x in xs) for h, xs in inv.items()}
    defectos = {f"{h}!{x['celda']}": x["defecto"] for h, xs in inv.items() for x in xs if x["defecto"] is not None}
    path = os.path.join(ROOT, "src", "data", "migracion.json")
    with open(path, "w", encoding="utf-8") as f:
        json.dump({"entradas": entradas, "defectos": defectos, "versiones": versiones, "renombres": RENAMES, "anclas": anclas(tpl)},
                  f, ensure_ascii=False, separators=(",", ":"))
    print(f"src/data/migracion.json: {os.path.getsize(path) // 1024} KB")


if __name__ == "__main__":
    main()
